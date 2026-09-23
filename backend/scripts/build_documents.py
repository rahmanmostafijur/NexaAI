"""Builds the demo knowledge base files from the Markdown sources.

Two documents are rendered as multi-page PDFs (so answers cite real page
numbers) and one as DOCX, exercising every extraction path of the ingestion
pipeline. `<!-- page -->` markers in the Markdown become PDF page breaks.
"""

from __future__ import annotations

import re
from pathlib import Path

import docx
from fpdf import FPDF

from app.core.config import BACKEND_DIR

SOURCE_DIR = BACKEND_DIR / "data" / "knowledge_base"

# (source markdown, output filename, format)
DOCUMENTS: list[tuple[str, str, str]] = [
    ("return_policy.md", "Return_Policy.md", "markdown"),
    ("refund_policy.md", "Refund_Policy.md", "markdown"),
    ("shipping_policy.md", "Shipping_Policy.md", "markdown"),
    ("warranty_policy.md", "Warranty_Policy.pdf", "pdf"),
    ("employee_handbook.md", "Employee_Handbook.pdf", "pdf"),
    ("membership_policy.md", "Customer_Membership_Policy.docx", "docx"),
    ("product_guidelines.md", "Product_Guidelines.md", "markdown"),
    ("company_faq.md", "Company_FAQ.md", "markdown"),
]

_PAGE_BREAK = "<!-- page -->"
_BOLD = re.compile(r"\*\*(.+?)\*\*")


def _pdf_safe(text: str) -> str:
    """Core PDF fonts are Latin-1 only."""
    replacements = {"৳": "BDT ", "–": "-", "—": "-", "’": "'", "“": '"', "”": '"', "…": "..."}
    for source, target in replacements.items():
        text = text.replace(source, target)
    return text.encode("latin-1", "replace").decode("latin-1")


def render_pdf(markdown: str, output: Path) -> None:
    pdf = FPDF(format="A4")
    pdf.set_auto_page_break(auto=True, margin=18)
    pdf.set_margins(20, 20, 20)
    for page in markdown.split(_PAGE_BREAK):
        pdf.add_page()
        for raw in page.strip().splitlines():
            line = _pdf_safe(raw.rstrip())
            if not line.strip():
                pdf.ln(3)
            elif line.startswith("# "):
                pdf.set_font("Helvetica", "B", 18)
                pdf.multi_cell(0, 10, line[2:], new_x="LMARGIN", new_y="NEXT")
                pdf.ln(2)
            elif line.startswith("## "):
                pdf.set_font("Helvetica", "B", 13)
                pdf.ln(2)
                pdf.multi_cell(0, 7, line[3:], new_x="LMARGIN", new_y="NEXT")
            else:
                pdf.set_font("Helvetica", "", 10.5)
                text = "- " + line[2:] if line.startswith("- ") else line
                pdf.multi_cell(0, 5.5, text, markdown=True, new_x="LMARGIN", new_y="NEXT")
    pdf.output(str(output))


def render_docx(markdown: str, output: Path) -> None:
    document = docx.Document()
    rows: list[list[str]] = []

    def flush_table() -> None:
        if not rows:
            return
        table = document.add_table(rows=len(rows), cols=len(rows[0]))
        table.style = "Table Grid"
        for r, row in enumerate(rows):
            for c, value in enumerate(row):
                table.cell(r, c).text = value
        rows.clear()

    for raw in markdown.replace(_PAGE_BREAK, "").splitlines():
        line = raw.strip()
        if line.startswith("|"):
            cells = [cell.strip() for cell in line.strip("|").split("|")]
            if not all(set(cell) <= set("-: ") for cell in cells):
                rows.append(cells)
            continue
        flush_table()
        if not line:
            continue
        if line.startswith("# "):
            document.add_heading(line[2:], level=0)
        elif line.startswith("## "):
            document.add_heading(line[3:], level=1)
        elif line.startswith("- "):
            paragraph = document.add_paragraph(style="List Bullet")
            _add_runs(paragraph, line[2:])
        else:
            _add_runs(document.add_paragraph(), line)
    flush_table()
    document.save(str(output))


def _add_runs(paragraph, text: str) -> None:
    parts = _BOLD.split(text)
    for index, part in enumerate(parts):
        if part:
            paragraph.add_run(part).bold = index % 2 == 1


def build_all(output_dir: Path) -> list[Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    paths = []
    for source, filename, fmt in DOCUMENTS:
        markdown = (SOURCE_DIR / source).read_text(encoding="utf-8")
        target = output_dir / filename
        if fmt == "pdf":
            render_pdf(markdown, target)
        elif fmt == "docx":
            render_docx(markdown, target)
        else:
            target.write_text(markdown, encoding="utf-8")
        paths.append(target)
    return paths


if __name__ == "__main__":
    for path in build_all(BACKEND_DIR / "storage" / "seed_documents"):
        print(path)
